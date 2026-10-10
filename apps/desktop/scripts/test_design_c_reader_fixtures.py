# SPDX-License-Identifier: MIT
"""Original-byte checks only; these are not installed Reader acceptance."""
from contextlib import closing
import hashlib
import importlib.util
from io import BytesIO
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from zipfile import ZipFile

spec = importlib.util.spec_from_file_location("reader_fixtures", Path(__file__).with_name("design-c-reader-fixtures.py"))
fixtures = importlib.util.module_from_spec(spec)
spec.loader.exec_module(fixtures)


class ReaderOriginals(unittest.TestCase):
    def setUp(self):
        self.manifest, records = fixtures.fixture_set()
        self.data = {row["filename"]: data for row, data in records}

    def test_deterministic_bounded_eight_categories_and_hashes(self):
        self.assertEqual((self.manifest, self.data), (fixtures.fixture_set()[0], {
            row["filename"]: data for row, data in fixtures.fixture_set()[1]}))
        self.assertEqual({row["gate"] for row in self.manifest["files"]}, {
            "reader-" + name for name in ("pdf", "docx", "csv", "xlsx", "markdown", "html", "xml", "images")})
        self.assertEqual(len(self.data), 11)
        for row in self.manifest["files"]:
            data = self.data[row["filename"]]
            self.assertEqual(hashlib.sha256(data).hexdigest(), row["sha256"])
            self.assertEqual(len(data), row["bytes"])
            self.assertLess(len(data), 32000)
        for name in ("synthetic-reading.docx", "synthetic-reading.xlsx"):
            with ZipFile(BytesIO(self.data[name])) as archive:
                self.assertTrue(all(info.date_time == (2000, 1, 1, 0, 0, 0) for info in archive.infolist()))
                self.assertFalse(any(b'TargetMode="External"' in archive.read(info) for info in archive.infolist()))

    def test_pdf_real_parser_reads_two_distinct_physical_pages(self):
        import pypdfium2 as pdfium
        with pdfium.PdfDocument(self.data["synthetic-reading.pdf"]) as document:
            self.assertEqual(len(document), 2)
            for number, topic in enumerate(("CKD", "dialysis")):
                with closing(document[number]) as page, closing(page.get_textpage()) as text:
                    self.assertIn("Synthetic " + topic + " reading page", text.get_text_range())
                    bitmap = page.render(scale=1)
                    try:
                        self.assertEqual((bitmap.width, bitmap.height), (360, 240))
                    finally:
                        bitmap.close()

    def test_docx_real_parser_reads_paragraph_table_and_embedded_original(self):
        from docx import Document
        document = Document(BytesIO(self.data["synthetic-reading.docx"]))
        self.assertEqual(document.paragraphs[0].text, "Synthetic transplantation reading paragraph.")
        self.assertEqual(document.tables[0].cell(1, 1).text, "Synthetic row")
        self.assertEqual(document.part.rels["rId1"].target_part.blob, self.data["synthetic-reading.png"])
        self.assertEqual(len(document.inline_shapes), 1)

    def test_xlsx_real_parser_retains_sheets_cached_formula_and_number_format(self):
        from openpyxl import load_workbook
        data = self.data["synthetic-reading.xlsx"]
        formulas = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        values = load_workbook(BytesIO(data), read_only=True, data_only=True, keep_links=False)
        try:
            self.assertEqual(formulas.sheetnames, ["CKD", "Dialysis"])
            self.assertEqual(formulas["CKD"]["B2"].value, "=A2*2")
            self.assertEqual(values["CKD"]["B2"].value, 24)
            self.assertEqual(values["Dialysis"]["A2"].value, 0.25)
            self.assertEqual(formulas["Dialysis"]["A2"].number_format, "0.00%")
            self.assertIsNone(values["Dialysis"]["B2"].value)
        finally:
            formulas.close()
            values.close()

    def test_images_decode_actual_png_jpeg_and_distinct_tiff_frames(self):
        from PIL import Image
        for suffix, fmt in (("png", "PNG"), ("jpg", "JPEG"), ("tiff", "TIFF")):
            with Image.open(BytesIO(self.data["synthetic-reading." + suffix])) as image:
                self.assertEqual(image.format, fmt)
                self.assertEqual(image.size, (48, 32))
                self.assertEqual(image.n_frames if hasattr(image, "n_frames") else 1, 2 if fmt == "TIFF" else 1)
                if fmt == "TIFF":
                    first = image.getpixel((0, 0))
                    image.seek(1)
                    self.assertNotEqual(first, image.getpixel((0, 0)))

    def test_text_fixtures_have_real_table_window_and_active_denial_inputs(self):
        import csv
        from defusedxml.ElementTree import fromstring
        from defusedxml.common import DTDForbidden
        rows = list(csv.reader(self.data["synthetic-reading.csv"].decode().splitlines()))
        self.assertEqual(len(rows), 106)
        self.assertEqual(rows[100][1], "Synthetic row 100, quoted")
        for suffix in ("html", "xml"):
            text = self.data["synthetic-reading." + suffix].decode()
            self.assertIn(fixtures.MARKER, text)
            self.assertIn(fixtures.EXTERNAL, text)
        self.assertEqual(fromstring(self.data["synthetic-reading.xml"], forbid_dtd=True).tag, "reading")
        with self.assertRaises(DTDForbidden):
            fromstring(self.data["synthetic-denied-entity.xml"], forbid_dtd=True, forbid_entities=True, forbid_external=True)

    def test_writer_requires_fresh_absolute_output_and_keeps_originals_immutable(self):
        with TemporaryDirectory(prefix="dc21-reader-unit-") as directory:
            output = Path(directory) / "originals"
            fixtures.write_fixtures(output)
            self.assertEqual(json.loads((output / "manifest.json").read_text(encoding="utf-8")), self.manifest)
            with self.assertRaises(FileExistsError):
                fixtures.write_fixtures(output)
            for name, data in self.data.items():
                self.assertEqual((output / name).read_bytes(), data)
        with self.assertRaisesRegex(ValueError, "absolute"):
            fixtures.write_fixtures(Path("relative-fixtures"))


if __name__ == "__main__":
    unittest.main()
