# SPDX-License-Identifier: MIT
"""Actual installed parser checks. Fixtures contain synthetic study material."""
from io import BytesIO
import json

import pytest

from renulus.contracts import ApiError
from renulus.knowledge.reader_binary import binary_asset, docx_payload, image_payload, pdf_payload, xlsx_payload
from renulus.knowledge.reader_models import ReaderRequest
from renulus.knowledge.reader_text import text_payload


def request(**kwargs):
    return ReaderRequest(reference={"target": {"type": "library_revision", "id": "synthetic_revision"}}, **kwargs)


def synthetic_pdf(text="SYNTHETIC dialysis access study"):
    stream = f"BT /F1 12 Tf 20 60 Td ({text}) Tj ET".encode("ascii")
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 240 100] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"]
    data, offsets = b"%PDF-1.4\n", [0]
    for number, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data += str(number).encode() + b" 0 obj\n" + obj + b"\nendobj\n"
    position = len(data)
    data += b"xref\n0 6\n0000000000 65535 f \n" + b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:])
    return data + b"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n" + str(position).encode() + b"\n%%EOF\n"


def test_actual_pdf_parser_text_and_png_raster_reject_signature_only_file():
    data = synthetic_pdf()
    payload = pdf_payload(data, request())
    assert payload.pages == 1 and payload.page == 1
    assert "dialysis access" in payload.blocks[0].text
    assert binary_asset(data, ".pdf", "pdf:1").startswith(b"\x89PNG\r\n\x1a\n")
    with pytest.raises(ApiError, match="PDF"):
        pdf_payload(b"%PDF-x\nSYNTHETIC invalid body", request())
    with pytest.raises(ApiError) as caught:
        binary_asset(data, ".pdf", "https://example.invalid/document")
    assert caught.value.code == "reader_asset_missing"


def test_actual_docx_headings_table_image_and_source_text_locations():
    from docx import Document
    from PIL import Image
    document = Document()
    document.add_heading("SYNTHETIC transplant study", 1)
    document.add_paragraph("SYNTHETIC rejection study paragraph")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "SYNTHETIC marker", "SYNTHETIC value"
    image = BytesIO(); Image.new("RGB", (20, 20), "white").save(image, format="PNG")
    document.add_picture(BytesIO(image.getvalue()))
    data = BytesIO(); document.save(data)
    payload = docx_payload(data.getvalue(), request())
    assert payload.blocks[0].kind == "heading" and payload.blocks[0].level == 1
    grid = next(item for item in payload.blocks if item.kind == "table")
    assert grid.grid.rows == (("SYNTHETIC marker", "SYNTHETIC value"),)
    picture = next(item for item in payload.blocks if item.kind == "image")
    assert binary_asset(data.getvalue(), ".docx", picture.asset_id).startswith(b"\x89PNG")
    assert all(item.locator.kind == "text" for item in payload.blocks if item.text)


def test_actual_xlsx_preserves_percent_zero_padding_literal_and_missing_formula_cache():
    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active; sheet.title = "SYNTHETIC CKD"
    sheet.append([0.15, 7, "=1+1", "SYNTHETIC literal"])
    sheet["A1"].number_format, sheet["B1"].number_format = "0%", "0000"
    workbook.create_sheet("SYNTHETIC dialysis")
    data = BytesIO(); workbook.save(data)
    payload = xlsx_payload(data.getvalue(), request())
    row = payload.table.rows[0]
    assert [cell.display for cell in row[:2]] == ["15%", "0007"]
    assert row[0].value == 0.15 and row[1].value == 7
    assert not row[2].cached and row[2].formula == "=1+1" and row[2].value is None
    assert "formula_cache_missing" in {item.code for item in payload.diagnostics}
    assert xlsx_payload(data.getvalue(), request(sheet="SYNTHETIC dialysis")).table.sheet_name == "SYNTHETIC dialysis"


def test_csv_literal_formulas_quoted_newlines_semicolon_pagination_and_bom():
    data = '\ufefftopic;value\n"SYNTHETIC\nAKI";=1+1\nSYNTHETIC transplant;0007\n'.encode()
    payload = text_payload(data, ".csv", request(first_row=2, row_limit=1))
    assert payload.table.total_rows == 3
    assert payload.table.first_row == 2
    assert [cell.display for cell in payload.table.rows[0]] == ["SYNTHETIC\nAKI", "=1+1"]
    assert payload.table.rows[0][1].formula is None
    assert not text_payload(b"", ".csv", request()).table.rows


def test_markup_never_exports_active_elements_resource_urls_or_entity_authority():
    hostile = b'''<html><head><meta http-equiv="refresh" content="0;url=/api/v1/cases/saved"></head>
    <script>fetch('/api/v1/cases/saved')</script><iframe src="/api/v1/health"></iframe>
    <img src="file:///C:/private" onerror="fetch('/api/v1/health')">
    <h1>SYNTHETIC glomerular study</h1><p>SYNTHETIC readable paragraph</p>
    <svg><script>SENTINEL_DO_NOT_RENDER</script></svg></html>'''
    payload = text_payload(hostile, ".html", request())
    encoded = payload.model_dump_json()
    assert "SYNTHETIC glomerular study" in encoded and "readable paragraph" in encoded
    for forbidden in ("fetch(", "file:///", "/api/v1/", "SENTINEL_DO_NOT_RENDER", "src="):
        assert forbidden not in encoded
    with pytest.raises(ApiError) as caught:
        text_payload(b'<!DOCTYPE x [<!ENTITY ex SYSTEM "file:///C:/Renulus/scratch/no-read">]><x>&ex;</x>', ".xml", request())
    assert caught.value.code == "reader_xml_invalid"
    xml = text_payload(b'<study topic="SYNTHETIC"><text>Dialysis &amp; access</text></study>', ".xml", request())
    assert xml.blocks[0].kind == "code" and "&amp;" in xml.blocks[0].text
    markdown = text_payload(b'# SYNTHETIC CKD\n\n- Study one\n\n<script>literal only</script>\n\n[Source](https://example.invalid)', ".md", request())
    assert markdown.blocks[0].kind == "heading"
    assert any(item.kind == "list_item" for item in markdown.blocks)
    assert "https://example.invalid" not in markdown.model_dump_json()


def test_actual_multiframe_tiff_png_jpeg_are_inert_png_assets():
    from PIL import Image
    image = Image.new("RGB", (20, 20), "white")
    for format, suffix in (("PNG", ".png"), ("JPEG", ".jpg"), ("TIFF", ".tiff")):
        buffer = BytesIO()
        image.save(buffer, format=format, save_all=format == "TIFF", append_images=[image] if format == "TIFF" else [])
        payload = image_payload(buffer.getvalue(), request(page=2 if format == "TIFF" else 1))
        assert payload.pages == (2 if format == "TIFF" else 1)
        assert binary_asset(buffer.getvalue(), suffix, payload.asset_id).startswith(b"\x89PNG")
        assert not payload.blocks


def test_declared_limits_reject_csv_and_deep_xml_before_renderer_publication():
    with pytest.raises(ApiError) as caught:
        text_payload((",".join("x" for _ in range(201))).encode(), ".csv", request())
    assert caught.value.code == "reader_table_limit"
    with pytest.raises(ApiError) as caught:
        text_payload(("<x>" * 81 + "SYNTHETIC" + "</x>" * 81).encode(), ".xml", request())
    assert caught.value.code == "reader_markup_limit"


def test_new_markup_adapters_feed_actual_hybrid_chunker_with_synthetic_tokenizer(tmp_path):
    # Real installed Docling transform/serialization; not an embedding-helper proof.
    from tokenizers import Tokenizer, models, pre_tokenizers
    from renulus.knowledge.engines import DoclingExtractor
    tokenizer = Tokenizer(models.WordLevel({"[UNK]": 0}, unk_token="[UNK]"))
    tokenizer.pre_tokenizer = pre_tokenizers.Whitespace()
    extractor = DoclingExtractor(None)
    extractor._tokenizer = tokenizer
    samples = {".csv": b"topic,value\nSYNTHETIC AKI,15%\n", ".html": b"<p>SYNTHETIC transplant study</p><script>OMIT_SCRIPT</script>",
               ".xml": b"<study>SYNTHETIC glomerular context</study>"}
    for suffix, raw in samples.items():
        path = tmp_path / ("SYNTHETIC" + suffix); path.write_bytes(raw)
        extracted = extractor.extract_file(path, "SYNTHETIC extraction")
        assert extracted.passages and extracted.document["schema_name"] == "DoclingDocument"
        assert "OMIT_SCRIPT" not in json.dumps(extracted.document)
        locators = [location for passage in extracted.passages for location in passage["locators"]]
        assert locators and all("content_locator" in location for location in locators)
        volatile = extractor.extract_bytes(raw, path.name, "SYNTHETIC volatile extraction")
        assert volatile.passages == extracted.passages
