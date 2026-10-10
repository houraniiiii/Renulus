# SPDX-License-Identifier: MIT
"""Actual conversation retention, Knowledge003 persistence and reader calls."""
from concurrent.futures import ThreadPoolExecutor
from contextlib import closing, contextmanager
from io import BytesIO
import json
from pathlib import Path
import sqlite3
from threading import Event

from fastapi.testclient import TestClient
import pytest

from renulus.contracts import ApiError
from renulus.design_c_contracts import Identity
from renulus.knowledge.attachment_models import AttachmentRemoval, AttachmentUpload
from renulus.knowledge.models import own_text_rights
from renulus.learn.conversation_models import NewConversation
from renulus.server import create_app
from renulus.storage.assessment_policy import advance_policy, policy_snapshot
from renulus.storage import attachment_cleanup
from renulus.storage.invalidation import bump_context_epoch, context_epoch_in


@pytest.fixture
def app(tmp_path, monkeypatch):
    import socket
    connect = socket.socket.connect
    def forbidden(instance, address):
        if isinstance(address, tuple) and address[0] in {"127.0.0.1", "::1"}:
            return connect(instance, address)
        raise AssertionError("Attachment tests must not use an external network or provider")
    monkeypatch.setattr(socket.socket, "connect", forbidden)
    result = create_app(tmp_path / "synthetic-ordinary-attachments")
    try:
        yield result
    finally:
        result.state.services.get("knowledge").index.close()


def conversation(app, *, temporary=False, kind="learning"):
    return app.state.services.get("conversations").create(NewConversation(
        kind=kind, temporary=temporary, title="SYNTHETIC attachment chat"))


def upload_request(app, chat, *, key="SYNTHETIC_upload_001", filename="SYNTHETIC.md", rights=None):
    with closing(app.state.services.db.connect()) as conn:
        return AttachmentUpload(conversation=chat["identity"], expected_conversation_revision=chat["revision"],
            expected_policy_revision=policy_snapshot(conn)["revision"], expected_context_epoch=context_epoch_in(conn),
            idempotency_key=key, filename=filename, title="SYNTHETIC study file", rights=rights or own_text_rights(), metadata={})


def read_body(receipt):
    return {"reference": receipt["reference"], "conversation": receipt["conversation"]}


@contextmanager
def busy_conversations(services):
    """A different thread owns the lock throughout real reader HTTP requests."""
    held, release = Event(), Event()
    def hold():
        with services.get("conversations")._lock:
            held.set()
            assert release.wait(30), "Reader waited for the volatile owner lock"
    with ThreadPoolExecutor(max_workers=1) as executor:
        future = executor.submit(hold)
        try:
            assert held.wait(5)
            yield
        finally:
            release.set()
            future.result(timeout=5)


@pytest.mark.parametrize("kind", ["learning", "case"])
@pytest.mark.parametrize("operation", ["resolve", "asset", "original"])
def test_saved_reader_ignores_unrelated_volatile_owner_contention(app, kind, operation):
    from PIL import Image
    stream = BytesIO()
    Image.new("RGB", (20, 12), "white").save(stream, format="PNG")
    data = stream.getvalue()
    services, client = app.state.services, TestClient(app)
    chat = conversation(app, kind=kind)
    request = upload_request(app, chat, filename="SYNTHETIC.png")
    uploaded = client.post("/api/v1/attachments", content=data,
        headers={"x-renulus-attachment-options": request.model_dump_json()})
    assert uploaded.status_code == 201, uploaded.text
    receipt = uploaded.json()
    body = read_body(receipt)
    if operation == "asset":
        body = {**body, "asset_id": "image:1", "original_sha256": receipt["original"]["sha256"]}
    with busy_conversations(services):
        response = client.post("/api/v1/readers/" + operation, json=body)
        assert response.status_code == 200, response.text
    assert response.headers["cache-control"] == "no-store"
    if operation == "resolve":
        result = response.json()
        assert result["preview_state"] == "ready" and result["payload"]["asset_id"] == "image:1"
        assert result["original"]["sha256"] == receipt["original"]["sha256"]
    elif operation == "asset":
        with Image.open(BytesIO(response.content)) as raster:
            assert raster.size == (20, 12)
    else:
        assert response.content == data


@pytest.mark.parametrize("kind", ["learning", "case"])
def test_temporary_reader_still_fails_closed_during_volatile_owner_contention(app, kind):
    services, client = app.state.services, TestClient(app)
    chat = conversation(app, kind=kind, temporary=True)
    receipt = services.get("knowledge_attachments").upload(upload_request(app, chat), b"SYNTHETIC volatile original")
    with busy_conversations(services):
        for operation in ("resolve", "original"):
            response = client.post("/api/v1/readers/" + operation, json=read_body(receipt))
            assert response.status_code == 409, response.text
            assert response.json()["error"]["code"] == "reader_original_unavailable"
            assert response.json()["error"]["retryable"] is True
    assert client.post("/api/v1/readers/original", json=read_body(receipt)).content == b"SYNTHETIC volatile original"


@pytest.mark.parametrize("kind", ["learning", "case"])
def test_saved_pdf_text_raster_and_original_work_with_busy_conversations(app, kind):
    from test_reader_formats import synthetic_pdf
    services, client = app.state.services, TestClient(app)
    chat = conversation(app, kind=kind)
    data = synthetic_pdf("SYNTHETIC retained PDF under unrelated contention")
    request = upload_request(app, chat, filename="SYNTHETIC.pdf")
    uploaded = client.post("/api/v1/attachments", content=data,
        headers={"x-renulus-attachment-options": request.model_dump_json()})
    assert uploaded.status_code == 201, uploaded.text
    receipt = uploaded.json()
    with busy_conversations(services):
        resolved = client.post("/api/v1/readers/resolve", json=read_body(receipt))
        assert resolved.status_code == 200, resolved.text
        reader = resolved.json()
        assert reader["preview_state"] == "ready"
        assert reader["payload"]["blocks"][0]["text"] == "SYNTHETIC retained PDF under unrelated contention"
        raster = client.post("/api/v1/readers/asset", json={**read_body(receipt),
            "asset_id": reader["payload"]["asset_id"], "original_sha256": reader["original"]["sha256"]})
        assert raster.status_code == 200 and raster.content.startswith(b"\x89PNG\r\n\x1a\n"), raster.text[:200]
        original = client.post("/api/v1/readers/original", json=read_body(receipt))
        assert original.status_code == 200 and original.content == data


@pytest.mark.parametrize("kind", ["learning", "case"])
def test_saved_reader_keeps_context_integrity_and_real_test_policy_guards_under_contention(app, kind):
    services, client = app.state.services, TestClient(app)
    chat = conversation(app, kind=kind)
    data = b"SYNTHETIC immutable original"
    receipt = services.get("knowledge_attachments").upload(upload_request(app, chat), data)
    body = read_body(receipt)
    path = services.paths.library / "attachments" / receipt["id"] / "original.md"
    with busy_conversations(services):
        mismatch = client.post("/api/v1/readers/original", json={**body,
            "conversation": {**chat["identity"], "id": "SYNTHETIC_different_chat"}})
        assert mismatch.status_code == 403 and mismatch.json()["error"]["code"] == "reader_context_mismatch"
        try:
            path.write_bytes(b"X" * len(data))  # Same-size bytes must still fail SHA verification.
            damaged = client.post("/api/v1/readers/original", json=body)
            assert damaged.status_code == 409 and damaged.json()["error"]["code"] == "reader_original_integrity"
        finally:
            path.write_bytes(data)
        created = client.post("/api/v1/assessment/tests/reviewed", json={
            "idempotency_key": "SYNTHETIC_reader_policy_test", "count": 1})
        assert created.status_code == 200, created.text
        started = client.post("/api/v1/assessment/test-attempts", json={
            "version": {"type": "saved_test_version", "id": created.json()["version_id"]}, "mode": "test"},
            headers={"Idempotency-Key": "SYNTHETIC_reader_policy_start"})
        assert started.status_code == 201, started.text
        for operation in ("resolve", "original"):
            restricted = client.post("/api/v1/readers/" + operation, json=body)
            assert restricted.status_code == 409 and restricted.json()["error"]["code"] == "assistance_restricted"
        with services.db.transaction() as conn:
            assert services.get("knowledge_attachments").authorize_in(conn, receipt["id"],
                context=Identity.model_validate(chat["identity"])) == {"allowed": False, "reason": "restricted"}


@pytest.mark.parametrize("kind", ["learning", "case"])
def test_promotion_committed_authority_precedes_volatile_cleanup(app, monkeypatch, kind):
    from renulus.knowledge.attachments import AttachmentPromotion
    services, client = app.state.services, TestClient(app)
    chat = conversation(app, kind=kind, temporary=True)
    attachment_owner = services.get("knowledge_attachments")
    data = b"SYNTHETIC promoted original"
    receipt = attachment_owner.upload(upload_request(app, chat), data)
    real_commit = AttachmentPromotion.commit
    observed = []
    def verify_committed(lease):
        # The real promotion holds Chat+Case locks after SQLite committed, but
        # before it drops either volatile attachment bytes or the temporary chat.
        assert receipt["id"] in attachment_owner._volatile
        def read_in_other_thread():
            with services.db.transaction() as conn:
                authorization = attachment_owner.authorize_in(conn, receipt["id"],
                    context=Identity.model_validate(chat["identity"]))
            volatile = attachment_owner.volatile_original(receipt["id"], context=Identity.model_validate(chat["identity"]))
            return authorization, volatile
        with ThreadPoolExecutor(max_workers=1) as executor:
            observed.append(executor.submit(read_in_other_thread).result(timeout=5))
        real_commit(lease)
    monkeypatch.setattr(AttachmentPromotion, "commit", verify_committed)
    saved = client.post(f"/api/v1/conversations/{kind}/{chat['identity']['id']}/save",
        json={"expected_revision": receipt["conversation_revision"]},
        headers={"Idempotency-Key": "SYNTHETIC_promote_lock_boundary"})
    assert saved.status_code == 200 and saved.json()["retention"] == "saved", saved.text
    assert observed == [({"allowed": True}, None)]
    assert receipt["id"] not in attachment_owner._volatile
    assert client.post("/api/v1/readers/original", json=read_body(receipt)).content == data


@pytest.mark.parametrize("kind", ["learning", "case"])
def test_uncommitted_promotion_is_not_saved_authority_and_rollback_keeps_original(app, monkeypatch, kind):
    from renulus.learn.conversation_models import Mutation
    services = app.state.services
    chat = conversation(app, kind=kind, temporary=True)
    owner = services.get("conversations")
    access = services.get("knowledge_attachments")
    data = b"SYNTHETIC uncommitted attachment"
    receipt = access.upload(upload_request(app, chat), data)
    original_insert = owner._insert_sidecar
    def fail_after_real_insert(conn, *args):
        original_insert(conn, *args)
        assert conn.execute("SELECT id FROM knowledge_attachments WHERE id=?", (receipt["id"],)).fetchone()
        def check_other_snapshot():
            with closing(services.db.connect()) as observer:
                observer.execute("BEGIN")
                assert not observer.execute("SELECT id FROM knowledge_attachments WHERE id=?", (receipt["id"],)).fetchone()
                return access.authorize_in(observer, receipt["id"], context=Identity.model_validate(chat["identity"]))
        with ThreadPoolExecutor(max_workers=1) as executor:
            assert executor.submit(check_other_snapshot).result(timeout=5) == {"allowed": False, "reason": "unavailable"}
        raise sqlite3.OperationalError("SYNTHETIC promotion rollback")
    with monkeypatch.context() as patch:
        patch.setattr(owner, "_insert_sidecar", fail_after_real_insert)
        with pytest.raises(sqlite3.OperationalError, match="SYNTHETIC promotion rollback"):
            owner.promote(kind, chat["identity"]["id"], Mutation(expected_revision=receipt["conversation_revision"]),
                "SYNTHETIC_promotion_rollback")
    assert not services.db.fetch_one("SELECT id FROM knowledge_attachments WHERE id=?", (receipt["id"],))
    assert receipt["id"] in access._volatile
    assert not (services.paths.library / "attachments" / receipt["id"]).exists()
    assert TestClient(app).post("/api/v1/readers/original", json=read_body(receipt)).content == data


@pytest.mark.parametrize("change,code", [("rights", "reader_context_mismatch"),
    ("delete", "reader_content_missing"), ("epoch", "reader_revision_conflict"),
    ("policy", "assistance_changed")])
def test_reader_rechecks_canonical_boundary_after_real_parse(app, monkeypatch, change, code):
    services, client = app.state.services, TestClient(app)
    chat = conversation(app)
    upload = upload_request(app, chat)
    receipt = services.get("knowledge_attachments").upload(upload, b"SYNTHETIC do not return after change")
    reader = services.get("readers")
    real_payload = reader.payload
    def change_after_parse(original, request):
        payload = real_payload(original, request)
        if change == "delete":
            services.get("knowledge_attachments").remove(receipt["id"], AttachmentRemoval(
                conversation=chat["identity"], expected_conversation_revision=receipt["conversation_revision"],
                expected_attachment_revision=1, expected_policy_revision=upload.expected_policy_revision,
                expected_context_epoch=upload.expected_context_epoch, idempotency_key="SYNTHETIC_delete_during_parse"))
        else:
            with services.db.transaction() as conn:
                if change == "rights":
                    rights = own_text_rights().model_copy(update={"display": False})
                    conn.execute("UPDATE knowledge_attachment_metadata SET rights_json=? WHERE attachment_id=?",
                        (rights.model_dump_json(), receipt["id"]))
                elif change == "epoch":
                    bump_context_epoch(conn)
                else:
                    advance_policy(conn)
        return payload
    monkeypatch.setattr(reader, "payload", change_after_parse)
    result = client.post("/api/v1/readers/resolve", json=read_body(receipt))
    assert result.status_code in {403, 404, 409} and result.json()["error"]["code"] == code, result.text
    assert "SYNTHETIC do not return" not in result.text


def test_inventory_supplies_actual_upload_pins_and_stale_restore_pin_fails(app):
    services = app.state.services
    chat = conversation(app)
    client = TestClient(app)
    path = "/api/v1/attachments/learning/" + chat["identity"]["id"]
    response = client.get(path)
    assert response.status_code == 200, response.text
    inventory = response.json()
    assert inventory["items"] == []
    options = {"conversation": inventory["conversation"],
        "expected_conversation_revision": inventory["conversation_revision"],
        "expected_policy_revision": inventory["policy_revision"],
        "expected_context_epoch": inventory["context_epoch"],
        "idempotency_key": "SYNTHETIC_inventory_upload", "filename": "SYNTHETIC.md",
        "title": "SYNTHETIC selected file", "rights": own_text_rights().model_dump(), "metadata": {}}
    uploaded = client.post("/api/v1/attachments", content=b"SYNTHETIC Attach selection",
        headers={"x-renulus-attachment-options": json.dumps(options)})
    assert uploaded.status_code == 201, uploaded.text
    receipt = uploaded.json()
    assert receipt["conversation_revision"] == inventory["conversation_revision"] + 1
    assert receipt["reference"] == {"target": {"type": "attachment", "id": receipt["id"]}, "locator": None}
    refreshed = client.get(path).json()
    options["expected_conversation_revision"] = refreshed["conversation_revision"]
    options["idempotency_key"] = "SYNTHETIC_inventory_stale"
    with services.db.transaction() as conn:
        bump_context_epoch(conn)
    stale = client.post("/api/v1/attachments", content=b"SYNTHETIC second selection",
        headers={"x-renulus-attachment-options": json.dumps(options)})
    assert stale.status_code == 409 and stale.json()["error"]["code"] == "attachment_context_changed"
    assert len(client.get(path).json()["items"]) == 1


def test_saved_upload_persists_exact_original_sidecar_and_reopens_without_extraction(app):
    services = app.state.services
    chat = conversation(app)
    request = upload_request(app, chat)
    client = TestClient(app)
    data = b"# SYNTHETIC AKI\n\nStudy reversible causes."
    response = client.post("/api/v1/attachments", content=data,
        headers={"x-renulus-attachment-options": request.model_dump_json()})
    assert response.status_code == 201, response.text
    receipt = response.json()
    assert receipt["retention"] == "saved" and receipt["conversation_revision"] == chat["revision"] + 1
    assert receipt["extraction_state"] == receipt["index_state"] == "not_requested"
    assert "original_path" not in response.text and "data" not in receipt
    original = services.db.fetch_one("SELECT * FROM knowledge_attachment_originals WHERE attachment_id=?", (receipt["id"],))
    expected = services.paths.library / "attachments" / receipt["id"] / "original.md"
    assert Path(original["original_path"]) == expected and expected.read_bytes() == data
    metadata = services.db.fetch_one("SELECT * FROM knowledge_attachment_metadata WHERE attachment_id=?", (receipt["id"],))
    assert metadata["filename"] == "SYNTHETIC.md" and metadata["source_id"] == "R02"
    assert json.loads(metadata["metadata_json"])["original_sha256"] == receipt["original"]["sha256"]
    assert metadata["preview_json"] is None and metadata["extraction_json"] is None
    result = client.post("/api/v1/readers/resolve", json=read_body(receipt))
    assert result.status_code == 200 and result.json()["payload"]["blocks"][0]["text"] == "SYNTHETIC AKI"
    reader = result.json()
    draft_url = "/api/v1/conversations/learning/" + chat["identity"]["id"] + "/draft"
    prior = client.put(draft_url, json={"text": "SYNTHETIC prior draft", "content_refs": [],
        "expected_revision": receipt["conversation_revision"]})
    assert prior.status_code == 200
    item = reader["payload"]["blocks"][1]
    passage = client.post("/api/v1/readers/passage", json={**read_body(receipt), "locator": item["locator"],
        "original_sha256": reader["original"]["sha256"], "expected_policy_revision": reader["policy_revision"],
        "expected_context_epoch": reader["context_epoch"]})
    assert passage.status_code == 200 and passage.json()["text"] == "Study reversible causes."
    merged = client.put(draft_url, json={"text": prior.json()["text"] + "\n\n" + passage.json()["text"],
        "content_refs": [passage.json()["reference"]], "expected_revision": prior.json()["conversation_revision"]})
    assert merged.status_code == 200
    saved_draft = client.get(draft_url).json()
    assert saved_draft["text"] == "SYNTHETIC prior draft\n\nStudy reversible causes."
    assert saved_draft["content_refs"] == [passage.json()["reference"]]
    assert services.db.fetch_one("SELECT COUNT(*) AS count FROM learn_messages WHERE thread_id=?", (chat["identity"]["id"],))["count"] == 0
    services.db.execute("UPDATE knowledge_attachments SET extraction_state='failed',index_state='failed' WHERE id=?", (receipt["id"],))
    assert client.post("/api/v1/readers/original", json=read_body(receipt)).content == data
    reopened = create_app(services.paths.root)
    try:
        assert TestClient(reopened).post("/api/v1/readers/original", json=read_body(receipt)).content == data
    finally:
        reopened.state.services.get("knowledge").index.close()


@pytest.mark.parametrize("kind", ["learning", "case"])
def test_temporary_upload_keeps_bytes_and_commands_out_of_canonical_storage(app, kind):
    services = app.state.services
    chat = conversation(app, temporary=True, kind=kind)
    rights = own_text_rights().model_copy(update={"cache": False})
    request = upload_request(app, chat, rights=rights)
    data = b"SYNTHETIC temporary dialysis study"
    receipt = services.get("knowledge_attachments").upload(request, data)
    assert receipt["retention"] == "temporary"
    assert not services.db.fetch_one("SELECT 1 FROM knowledge_attachments WHERE id=?", (receipt["id"],))
    assert not services.db.fetch_one("SELECT 1 FROM conversation_commands WHERE idempotency_key=?", (request.idempotency_key,))
    assert not (services.paths.library / "attachments" / receipt["id"]).exists()
    client = TestClient(app)
    result = client.post("/api/v1/readers/resolve", json=read_body(receipt))
    assert result.status_code == 200 and result.json()["authority"] == "knowledge"
    assert client.post("/api/v1/readers/original", json=read_body(receipt)).content == data
    with services.db.transaction() as conn:
        result = services.get("knowledge_attachments").authorize_in(conn, receipt["id"],
            context=Identity.model_validate(chat["identity"]))
        assert result == {"allowed": True}
        assert services.get("knowledge_attachments").authorize_in(conn, receipt["id"]) == {
            "allowed": False, "reason": "not_found"}
    reopened = create_app(services.paths.root)
    try:
        assert TestClient(reopened).post("/api/v1/readers/resolve", json=read_body(receipt)).status_code == 404
    finally:
        reopened.state.services.get("knowledge").index.close()


@pytest.mark.parametrize("temporary", [False, True])
def test_retry_is_idempotent_and_different_bytes_or_revision_do_not_duplicate(app, temporary):
    services = app.state.services
    chat = conversation(app, temporary=temporary)
    request = upload_request(app, chat)
    owner = services.get("knowledge_attachments")
    receipt = owner.upload(request, b"SYNTHETIC original")
    assert owner.upload(request, b"SYNTHETIC original") == receipt
    assert services.get("conversations").get("learning", chat["identity"]["id"])["revision"] == receipt["conversation_revision"]
    with pytest.raises(ApiError) as caught:
        owner.upload(request, b"SYNTHETIC different")
    assert caught.value.code == "idempotency_conflict"
    stale = request.model_copy(update={"idempotency_key": "SYNTHETIC_stale_002"})
    with pytest.raises(ApiError) as caught:
        owner.upload(stale, b"SYNTHETIC second")
    assert caught.value.code == "conversation_revision_conflict"
    assert len(owner._volatile) == (1 if temporary else 0)
    rows = services.db.fetch_all("SELECT id FROM knowledge_attachments")
    assert len(rows) == (0 if temporary else 1)


def test_policy_and_context_changes_during_upload_refuse_publication(app, monkeypatch):
    services = app.state.services
    chat = conversation(app)
    request = upload_request(app, chat)
    import renulus.knowledge.attachments as module
    real = module.admit
    def changed(filename, data, **kwargs):
        result = real(filename, data, **kwargs)
        with services.db.transaction() as conn:
            bump_context_epoch(conn)
        return result
    monkeypatch.setattr(module, "admit", changed)
    with pytest.raises(ApiError) as caught:
        services.get("knowledge_attachments").upload(request, b"SYNTHETIC stale context")
    assert caught.value.code == "attachment_context_changed"
    assert not services.db.fetch_all("SELECT id FROM knowledge_attachments")
    assert not (services.paths.library / "attachments").exists()
    current = upload_request(app, chat)
    monkeypatch.setattr(module, "admit", real)
    with services.db.transaction() as conn:
        advance_policy(conn)
    with pytest.raises(ApiError) as caught:
        services.get("knowledge_attachments").upload(current, b"SYNTHETIC stale policy")
    assert caught.value.code == "assistance_changed"


def test_real_original_write_failure_leaves_no_attachment_or_conversation_revision(app, monkeypatch):
    services = app.state.services
    chat = conversation(app)
    import renulus.knowledge.attachments as module
    def failed(self, conn, record, path):
        assert path.read_bytes() == b"SYNTHETIC disk rollback"
        raise sqlite3.OperationalError("SYNTHETIC insert failure after original write")
    monkeypatch.setattr(module.KnowledgeAttachments, "_insert_in", failed)
    with pytest.raises(sqlite3.OperationalError):
        services.get("knowledge_attachments").upload(upload_request(app, chat), b"SYNTHETIC disk rollback")
    assert not services.db.fetch_all("SELECT id FROM knowledge_attachments")
    assert not list((services.paths.library / "attachments").iterdir())
    assert services.get("conversations").get("learning", chat["identity"]["id"])["revision"] == chat["revision"]


@pytest.mark.parametrize("temporary", [False, True])
def test_remove_revokes_reader_and_cleans_only_owned_original_with_safe_retry(app, tmp_path, temporary):
    services = app.state.services
    chat = conversation(app, temporary=temporary)
    upload = upload_request(app, chat)
    receipt = services.get("knowledge_attachments").upload(upload, b"SYNTHETIC removed study file")
    sentinel = services.paths.library / "SYNTHETIC-unowned.txt"
    sentinel.write_bytes(b"SYNTHETIC preserve")
    removal = AttachmentRemoval(conversation=chat["identity"], expected_conversation_revision=receipt["conversation_revision"],
        expected_policy_revision=upload.expected_policy_revision, expected_context_epoch=upload.expected_context_epoch,
        idempotency_key="SYNTHETIC_remove_001", expected_attachment_revision=1)
    owner = services.get("knowledge_attachments")
    result = owner.remove(receipt["id"], removal)
    assert owner.remove(receipt["id"], removal) == result
    assert sentinel.read_bytes() == b"SYNTHETIC preserve"
    assert not (services.paths.library / "attachments" / receipt["id"]).exists()
    assert TestClient(app).post("/api/v1/readers/original", json=read_body(receipt)).status_code == 404
    assert services.db.is_deleted("attachment", receipt["id"]) is not temporary
    assert services.db.is_deleted("attachment_original", receipt["original"]["id"]) is not temporary
    with pytest.raises(ApiError):
        owner.upload(upload, b"SYNTHETIC removed study file")


def test_saved_upload_requires_cache_and_raw_api_rejects_path_and_invalid_signature(app):
    chat = conversation(app)
    rights = own_text_rights().model_copy(update={"cache": False})
    request = upload_request(app, chat, rights=rights)
    client = TestClient(app)
    denied = client.post("/api/v1/attachments", content=b"SYNTHETIC no cache", headers={
        "x-renulus-attachment-options": request.model_dump_json()})
    assert denied.status_code == 403
    bad = request.model_dump(mode="json"); bad["filename"] = "../SYNTHETIC.md"
    assert client.post("/api/v1/attachments", content=b"SYNTHETIC", headers={
        "x-renulus-attachment-options": json.dumps(bad)}).status_code == 422
    pdf = upload_request(app, chat, filename="SYNTHETIC.pdf")
    assert client.post("/api/v1/attachments", content=b"SYNTHETIC not PDF", headers={
        "x-renulus-attachment-options": pdf.model_dump_json()}).status_code == 422
    assert not app.state.services.db.fetch_all("SELECT id FROM knowledge_attachments")


def test_delete_chat_plan_tombstones_processing_and_rolls_back_without_deleting_bytes(app, monkeypatch):
    services = app.state.services
    chat = conversation(app)
    owner = services.get("knowledge_attachments")
    receipt = owner.upload(upload_request(app, chat), b"SYNTHETIC delete transaction")
    path = services.paths.library / "attachments" / receipt["id"] / "original.md"
    services.db.execute("UPDATE knowledge_attachment_metadata SET processing_json=? WHERE attachment_id=?",
        (json.dumps({"schema_version": "SYNTHETIC_processing"}), receipt["id"]))
    conn = services.db.connect()
    try:
        conn.execute("BEGIN IMMEDIATE")
        def forbidden(*args, **kwargs):
            raise AssertionError("delete_chat_in cannot use files or nested connections")
        with monkeypatch.context() as patch:
            patch.setattr(services.db, "connect", forbidden)
            for method in ("resolve", "open", "stat", "exists", "is_file", "mkdir"):
                patch.setattr(Path, method, forbidden)
            plan = owner.delete_chat_in(conn, "learning", chat["identity"]["id"])
        assert plan == [{"attachment_id": receipt["id"], "relative_directory": "library/attachments/" + receipt["id"]}]
        assert path.read_bytes() == b"SYNTHETIC delete transaction"
        assert not conn.execute("SELECT 1 FROM knowledge_attachment_metadata WHERE attachment_id=?", (receipt["id"],)).fetchone()
        assert conn.execute("SELECT 1 FROM deletion_ledger WHERE entity_type='attachment_original' AND entity_id=?",
            (receipt["original"]["id"],)).fetchone()
    finally:
        conn.rollback(); conn.close()
    assert services.db.fetch_one("SELECT processing_json FROM knowledge_attachment_metadata WHERE attachment_id=?", (receipt["id"],))
    assert not services.db.is_deleted("attachment", receipt["id"])
    with services.db.transaction() as conn:
        committed = owner.delete_chat_in(conn, "learning", chat["identity"]["id"])
    assert committed == plan and path.exists()
    assert services.db.is_deleted("attachment", receipt["id"])
    assert TestClient(app).post("/api/v1/readers/original", json=read_body(receipt)).status_code == 404
    # Parent's live Delete file journal consumes the inert plan. This service
    # helper is post-commit cleanup only; the hook itself never deletes bytes.
    owner.cleanup_deleted([entry["attachment_id"] for entry in plan])
    assert not path.exists()


def test_volatile_forget_releases_payload_and_retry_receipts_without_durable_deletion(app):
    services = app.state.services
    chat = conversation(app, temporary=True)
    owner = services.get("knowledge_attachments")
    assert services.get("chat_attachments") is owner
    request = upload_request(app, chat)
    receipt = owner.upload(request, b"SYNTHETIC volatile forget")
    owner.discard_volatile("learning", chat["identity"]["id"])
    assert not owner._volatile and not owner._volatile_commands
    assert not services.db.fetch_one("SELECT 1 FROM deletion_ledger WHERE entity_id=?", (receipt["id"],))
    assert not (services.paths.library / "attachments").exists()
    assert TestClient(app).post("/api/v1/readers/original", json=read_body(receipt)).status_code == 404


def test_promotion_stage_preserves_volatile_bytes_until_caller_commit_or_rollback(app):
    services = app.state.services
    chat = conversation(app, temporary=True)
    owner, conversations = services.get("knowledge_attachments"), services.get("conversations")
    receipt = owner.upload(upload_request(app, chat), b"SYNTHETIC staged promotion")
    path = services.paths.library / "attachments" / receipt["id"] / "original.md"
    with owner.promotion_guard(), conversations._lock, conversations.cases._lock:
        conn = services.db.connect()
        try:
            conn.execute("BEGIN IMMEDIATE")
            lease = owner.prepare_promotion_in(conn, "learning", chat["identity"]["id"])
            assert path.exists() and receipt["id"] in owner._volatile
            assert conn.execute("SELECT 1 FROM knowledge_attachments WHERE id=?", (receipt["id"],)).fetchone()
        finally:
            conn.rollback(); conn.close()
        lease.rollback()
        assert not path.exists() and receipt["id"] in owner._volatile
        with services.db.transaction() as conn:
            lease = owner.prepare_promotion_in(conn, "learning", chat["identity"]["id"])
            # Exercise the actual caller-connection Learn owner, not a fabricated
            # chat table. Zeno still owns wiring this into Conversation.promote.
            conversations.learn.create_thread(chat["title"], conn=conn, identifier=chat["identity"]["id"])
            conversations._insert_sidecar(conn, "learning", chat["identity"]["id"])
        lease.commit()
        conversations._temporary.pop(("learning", chat["identity"]["id"]))
    assert receipt["id"] not in owner._volatile and path.read_bytes() == b"SYNTHETIC staged promotion"
    assert TestClient(app).post("/api/v1/readers/original", json=read_body(receipt)).content == b"SYNTHETIC staged promotion"


def test_promotion_refuses_ungranted_cache_before_any_file_write(app):
    services = app.state.services
    chat = conversation(app, temporary=True)
    owner = services.get("knowledge_attachments")
    owner.upload(upload_request(app, chat, rights=own_text_rights().model_copy(update={"cache": False})), b"SYNTHETIC cache denied")
    with owner.promotion_guard(), services.get("conversations")._lock, services.db.transaction() as conn:
        with pytest.raises(ApiError) as caught:
            owner.prepare_promotion_in(conn, "learning", chat["identity"]["id"])
        assert caught.value.code == "attachment_cache_denied"
    assert not (services.paths.library / "attachments").exists()


def test_ordinary_upload_all_required_formats_connect_to_actual_reader_and_native_assets(app):
    from docx import Document
    from openpyxl import Workbook
    from PIL import Image
    from test_reader_formats import synthetic_pdf
    docx = Document(); docx.add_heading("SYNTHETIC transplant study", 1)
    table = docx.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "SYNTHETIC marker", "0015"
    document = BytesIO(); docx.save(document)
    xlsx = Workbook(); xlsx.active.append(["SYNTHETIC topic", 0.15]); xlsx.active["B1"].number_format = "0%"
    workbook = BytesIO(); xlsx.save(workbook)
    files = {"SYNTHETIC.pdf": synthetic_pdf(), "SYNTHETIC.docx": document.getvalue(),
        "SYNTHETIC.csv": b"topic,value\nSYNTHETIC CKD,15%\n", "SYNTHETIC.xlsx": workbook.getvalue(),
        "SYNTHETIC.md": b"# SYNTHETIC AKI\n\n| Topic | Value |\n| --- | --- |\n| CKD | 15% |",
        "SYNTHETIC.html": b'<table><tr><td>SYNTHETIC dialysis</td></tr></table><script>fetch("/api/v1/health")</script>',
        "SYNTHETIC.xml": b'<study>SYNTHETIC glomerular study</study>'}
    for format, extension in (("PNG", "png"), ("JPEG", "jpg"), ("TIFF", "tiff")):
        image = BytesIO(); Image.new("RGB", (20, 20), "white").save(image, format=format)
        files["SYNTHETIC." + extension] = image.getvalue()
    services = app.state.services
    chat, client = conversation(app), TestClient(app)
    draft_url = "/api/v1/conversations/learning/" + chat["identity"]["id"] + "/draft"
    retained, draft_text, draft_refs = [], "SYNTHETIC retained draft", []
    for number, (filename, data) in enumerate(files.items()):
        request = upload_request(app, chat, key=f"SYNTHETIC_format_{number:03}", filename=filename)
        uploaded = client.post("/api/v1/attachments", content=data, headers={"x-renulus-attachment-options": request.model_dump_json()})
        assert uploaded.status_code == 201, uploaded.text
        receipt = uploaded.json()
        chat = services.get("conversations").get("learning", chat["identity"]["id"])
        result = client.post("/api/v1/readers/resolve", json=read_body(receipt))
        assert result.status_code == 200, result.text
        body = result.json()
        assert body["preview_state"] == "ready", (filename, body["error"])
        assert client.post("/api/v1/readers/original", json=read_body(receipt)).content == data
        assert "fetch(" not in json.dumps(body["payload"])
        retained.append((receipt, data, body["payload"]))
        if body["payload"]["asset_id"]:
            asset = client.post("/api/v1/readers/asset", json={**read_body(receipt),
                "asset_id": body["payload"]["asset_id"], "original_sha256": body["original"]["sha256"]})
            assert asset.status_code == 200 and asset.content.startswith(b"\x89PNG\r\n\x1a\n")
        if receipt["media_type"].startswith("image/"):
            assert body["capabilities"]["ocr"] is False and body["capabilities"]["passage"] is False
            # A decoded image preview must never attest to absent OCR text.
            locator = {"kind": "image", "frame": 1}
            selected = client.post("/api/v1/readers/passage", json={**read_body(receipt), "locator": locator,
                "original_sha256": body["original"]["sha256"], "expected_policy_revision": body["policy_revision"],
                "expected_context_epoch": body["context_epoch"]})
            assert selected.status_code == 422 and selected.json()["error"]["code"] == "reader_selection_unavailable"
            continue
        payload = body["payload"]
        if payload["table"]:
            grid = payload["table"]
            row_number = grid["first_row"] + len(grid["rows"]) - 1
            locator = {"kind": "table", "format": grid["format"], "sheet_name": grid["sheet_name"],
                "first_row": row_number, "last_row": row_number, "first_column": 1,
                "last_column": grid["total_columns"]}
            expected_text = "\t".join(cell["display"] for cell in grid["rows"][-1])
        else:
            item = next((item for item in payload["blocks"] if item["kind"] == "table"), payload["blocks"][0])
            locator, expected_text = item["locator"], item["text"]
        selected = client.post("/api/v1/readers/passage", json={**read_body(receipt), "locator": locator,
            "original_sha256": body["original"]["sha256"], "expected_policy_revision": body["policy_revision"],
            "expected_context_epoch": body["context_epoch"]})
        assert selected.status_code == 200, (filename, selected.text)
        passage = selected.json()
        assert passage["text"] == expected_text and passage["reference"]["locator"] == locator
        assert passage["reference"]["target"] == receipt["reference"]["target"]
        draft_text += "\n\n" + passage["text"]
        draft_refs.append(passage["reference"])
        merged = client.put(draft_url, json={"text": draft_text, "content_refs": draft_refs,
            "expected_revision": chat["revision"]})
        assert merged.status_code == 200, (filename, merged.text)
        chat = services.get("conversations").get("learning", chat["identity"]["id"])
        # Revoked model permission must remove passage authority, while the
        # same retained original and inert reading remain available.
        rights = own_text_rights().model_dump()
        rights["model_input"] = False
        services.db.execute("UPDATE knowledge_attachment_metadata SET rights_json=? WHERE attachment_id=?",
            (json.dumps(rights), receipt["id"]))
        restricted = client.post("/api/v1/readers/resolve", json=read_body(receipt))
        assert restricted.status_code == 200 and not restricted.json()["capabilities"]["passage"]
        assert client.post("/api/v1/readers/original", json=read_body(receipt)).content == data
        denied = client.post("/api/v1/readers/passage", json={**read_body(receipt), "locator": locator,
            "original_sha256": body["original"]["sha256"], "expected_policy_revision": body["policy_revision"],
            "expected_context_epoch": body["context_epoch"]})
        assert denied.status_code == 403 and denied.json()["error"]["code"] == "reader_passage_denied"
    listed = client.get("/api/v1/attachments/learning/" + chat["identity"]["id"])
    assert listed.status_code == 200 and len(listed.json()["items"]) == len(files)
    assert all(item["access"] == {"allowed": True} for item in listed.json()["items"])
    assert services.db.fetch_one("SELECT COUNT(*) AS count FROM learn_messages WHERE thread_id=?",
        (chat["identity"]["id"],))["count"] == 0
    reopened = create_app(services.paths.root)
    try:
        other = TestClient(reopened)
        for receipt, data, payload in retained:
            assert other.post("/api/v1/readers/original", json=read_body(receipt)).content == data
            result = other.post("/api/v1/readers/resolve", json=read_body(receipt))
            assert result.status_code == 200 and result.json()["payload"] == payload
            assert not result.json()["capabilities"]["passage"]
        draft = other.get(draft_url)
        assert draft.status_code == 200 and draft.json()["text"] == draft_text
        assert draft.json()["content_refs"] == draft_refs and len(draft_refs) == 7
    finally:
        reopened.state.services.get("knowledge").index.close()


def test_individual_delete_journals_locked_file_then_actual_startup_finishes_cleanup(app, monkeypatch):
    services = app.state.services
    chat = conversation(app)
    request = upload_request(app, chat)
    owner = services.get("chat_attachments")
    receipt = owner.upload(request, b"SYNTHETIC durable delete")
    path = services.paths.library / "attachments" / receipt["id"] / "original.md"
    removal = AttachmentRemoval(conversation=chat["identity"], expected_conversation_revision=receipt["conversation_revision"],
        expected_policy_revision=request.expected_policy_revision, expected_context_epoch=request.expected_context_epoch,
        idempotency_key="SYNTHETIC_journal_delete", expected_attachment_revision=1)
    def locked(*args):
        raise OSError("SYNTHETIC locked original")
    with monkeypatch.context() as patch:
        patch.setattr(attachment_cleanup, "remove_owned_tree", locked)
        deleted = owner.remove(receipt["id"], removal)
    assert deleted["deleted"] and path.exists()
    assert not services.db.fetch_one("SELECT 1 FROM knowledge_attachments WHERE id=?", (receipt["id"],))
    journal = services.db.fetch_one("SELECT value FROM preferences WHERE key=?", (attachment_cleanup.JOURNAL,))
    assert json.loads(journal["value"]) == [{"attachment_id": receipt["id"],
        "relative_directory": "library/attachments/" + receipt["id"]}]
    assert TestClient(app).post("/api/v1/readers/original", json=read_body(receipt)).status_code == 404
    reopened = create_app(services.paths.root)
    with TestClient(reopened):
        assert not path.exists()
    assert attachment_cleanup.drain(reopened.state.services) == {"removed": 0, "pending": 0}
    assert owner.remove(receipt["id"], removal) == deleted


def test_individual_delete_enqueue_failure_rolls_back_rows_tombstones_revision_and_bytes(app, monkeypatch):
    services = app.state.services
    chat = conversation(app)
    upload = upload_request(app, chat)
    receipt = services.get("chat_attachments").upload(upload, b"SYNTHETIC journal rollback")
    path = services.paths.library / "attachments" / receipt["id"] / "original.md"
    removal = AttachmentRemoval(conversation=chat["identity"], expected_conversation_revision=receipt["conversation_revision"],
        expected_policy_revision=upload.expected_policy_revision, expected_context_epoch=upload.expected_context_epoch,
        idempotency_key="SYNTHETIC_enqueue_failure", expected_attachment_revision=1)
    def failed(conn, plan):
        assert conn.in_transaction
        assert plan[0]["attachment_id"] == receipt["id"]
        assert conn.execute("SELECT 1 FROM deletion_ledger WHERE entity_type='attachment' AND entity_id=?",
            (receipt["id"],)).fetchone()
        raise sqlite3.OperationalError("SYNTHETIC journal write failure")
    monkeypatch.setattr(attachment_cleanup, "enqueue_in", failed)
    with pytest.raises(sqlite3.OperationalError):
        services.get("chat_attachments").remove(receipt["id"], removal)
    assert path.read_bytes() == b"SYNTHETIC journal rollback"
    assert not services.db.is_deleted("attachment", receipt["id"])
    assert not services.db.is_deleted("attachment_original", receipt["original"]["id"])
    assert services.get("conversations").get("learning", chat["identity"]["id"])["revision"] == receipt["conversation_revision"]
    assert not services.db.fetch_one("SELECT 1 FROM conversation_commands WHERE idempotency_key=?", (removal.idempotency_key,))
    assert TestClient(app).post("/api/v1/readers/original", json=read_body(receipt)).content == b"SYNTHETIC journal rollback"


@pytest.mark.parametrize("version", [1, 2])
def test_producer_original_round_trips_through_actual_typed_archive_and_reader(app, version):
    from renulus.storage.recovery_files import remove_owned_tree
    services = app.state.services
    chat = conversation(app)
    data = b"# SYNTHETIC typed recovery\n\nPreserve exact original and source permissions."
    receipt = services.get("chat_attachments").upload(upload_request(app, chat), data)
    directory, manifest = services.get("data_recovery").backup(format_version=version)
    try:
        archive = (directory / "backup.zip").read_bytes()
    finally:
        remove_owned_tree(services.paths, directory)
    assert manifest["originals"] == [{"kind": "chat-attachment", "path": "library/attachments/" + receipt["id"] + "/original.md",
        "attachment_id": receipt["id"], "original_id": receipt["original"]["id"], "chat_kind": "learning",
        "chat_id": chat["identity"]["id"], "bytes": len(data), "sha256": receipt["original"]["sha256"]}]
    target_app = create_app(services.paths.root.parent / ("synthetic-restored-archive-" + str(version)))
    target = target_app.state.services
    recovery = target.get("data_recovery")
    token, stage = recovery.begin_preview()
    keep = False
    try:
        (stage / "input.zip").write_bytes(archive)
        checked = recovery.complete_preview(token, stage)
        keep = True
    finally:
        recovery.end_upload(stage, keep=keep)
    try:
        restored = recovery.restore_preview(checked["preview_token"], checked["exported_at"], True)
        assert restored["restored_originals"] == 1
        assert TestClient(target_app).post("/api/v1/readers/original", json=read_body(receipt)).content == data
        result = TestClient(target_app).post("/api/v1/readers/resolve", json=read_body(receipt))
        assert result.status_code == 200 and result.json()["payload"]["blocks"][0]["text"] == "SYNTHETIC typed recovery"
        metadata = target.db.fetch_one("SELECT * FROM knowledge_attachment_metadata WHERE attachment_id=?", (receipt["id"],))
        assert metadata["filename"] == "SYNTHETIC.md" and json.loads(metadata["rights_json"])["display"]
        assert not target.db.fetch_all("SELECT * FROM memory_jobs")
    finally:
        target.get("knowledge").index.close()
