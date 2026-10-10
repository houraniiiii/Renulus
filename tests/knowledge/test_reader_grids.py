# SPDX-License-Identifier: MIT
"""Actual format parsers preserve inert document tables and source coordinates."""
from io import BytesIO

import pytest

from renulus.contracts import ApiError
from renulus.knowledge.reader_binary import docx_payload
from renulus.knowledge.reader_models import ReaderRequest
from renulus.knowledge.reader_text import text_payload


def request():
    return ReaderRequest(reference={"target": {"type": "library_revision", "id": "SYNTHETIC_grid"}})


@pytest.mark.parametrize("suffix,content", [
    (".md", b"# SYNTHETIC study\n\n| Topic | Value |\n| --- | --- |\n| CKD | 15% |\n\nAfter the table."),
    (".html", b'<h1>SYNTHETIC study</h1><table><tr><th>Topic</th><th>Value</th></tr><tr><td>CKD</td><td>15%</td></tr></table><p>After the table.</p>')])
def test_markup_keeps_table_cells_order_headers_and_exact_text_locator(suffix, content):
    payload = text_payload(content, suffix, request())
    table = next(item for item in payload.blocks if item.kind == "table")
    assert table.grid.rows == (("Topic", "Value"), ("CKD", "15%"))
    assert table.grid.header_rows == 1
    assert table.text == "Topic\tValue\nCKD\t15%"
    assert table.locator.kind == "text" and table.locator.item_ref == table.item_ref
    assert table.locator.start == 0 and table.locator.end == len(table.text)
    assert payload.blocks[0].kind == "heading" and payload.blocks[-1].text == "After the table."


def test_docx_grid_preserves_actual_cells_between_surrounding_paragraphs():
    from docx import Document
    document = Document()
    document.add_paragraph("SYNTHETIC before")
    table = document.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, (("Topic", "Value"), ("Dialysis", "0015")), strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    document.add_paragraph("SYNTHETIC after")
    stream = BytesIO(); document.save(stream)
    payload = docx_payload(stream.getvalue(), request())
    grid = next(item for item in payload.blocks if item.kind == "table")
    assert grid.grid.rows == (("Topic", "Value"), ("Dialysis", "0015"))
    assert grid.grid.header_rows == 0  # Do not invent a header the source did not mark.
    assert grid.locator.end == len(grid.text)
    assert [item.kind for item in payload.blocks] == ["paragraph", "table", "paragraph"]


def test_html_cell_content_never_acquires_resource_or_script_authority():
    payload = text_payload(b'<table><tr><td>SYNTHETIC<img src="file:///C:/no-read" onerror="fetch()"><script>SENTINEL_OMIT</script></td><td><a href="/api/v1/private">Study source</a></td></tr></table>', ".html", request())
    table = payload.blocks[0]
    assert table.grid.rows == (("SYNTHETIC", "Study source"),)
    assert table.grid.header_rows == 0
    serialized = payload.model_dump_json()
    assert all(value not in serialized for value in ("file:///", "fetch()", "SENTINEL_OMIT", "/api/v1/private"))


def test_nested_or_over_limit_html_tables_fail_explicitly():
    with pytest.raises(ApiError) as caught:
        text_payload(b"<table><tr><td><table><tr><td>SYNTHETIC</td></tr></table></td></tr></table>", ".html", request())
    assert caught.value.code == "reader_table_unsupported"
    with pytest.raises(ApiError) as caught:
        text_payload(("<table>" + "<tr><td>SYNTHETIC</td></tr>" * 201 + "</table>").encode(), ".html", request())
    assert caught.value.code == "reader_table_limit"


def test_spanning_html_and_merged_docx_fail_without_inventing_cell_coordinates():
    with pytest.raises(ApiError) as caught:
        text_payload(b'<table><tr><td colspan="2">SYNTHETIC merged value</td></tr></table>', ".html", request())
    assert caught.value.code == "reader_table_unsupported"
    from docx import Document
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "SYNTHETIC merged value"
    stream = BytesIO(); document.save(stream)
    with pytest.raises(ApiError) as caught:
        docx_payload(stream.getvalue(), request())
    assert caught.value.code == "reader_table_unsupported"
