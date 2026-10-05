"""Actual pinned Docling and HybridChunker; no OCR/model or provider is loaded."""
import hashlib
import ipaddress
import json
import os
from pathlib import Path
import re
import socket
import subprocess
from io import BytesIO
from zipfile import ZipFile

import pytest

from renulus.contracts import ApiError
from renulus.knowledge.engines import CHUNK_TOKENS, DoclingExtractor
from renulus.knowledge.office import OFFICE_MEDIA


def make_office(path, *, long_table=False, title_only=False):
    if path.suffix == ".docx":
        from docx import Document
        document = Document()
        document.add_heading("Synthetic glomerular study", 0)
        if not title_only:
            document.add_paragraph("Synthetic kidney observation about hematuria.")
            table = document.add_table(rows=1, cols=2)
            table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Topic", "Observation"
            for n in range(70 if long_table else 2):
                cells = table.add_row().cells
                cells[0].text = "ROW" + str(n)
                cells[1].text = "synthetic glomerular proteinuria learning " * (12 if long_table else 1)
        document.save(path)
    elif path.suffix == ".pptx":
        from pptx import Presentation
        from pptx.util import Inches
        presentation = Presentation()
        for topic in ("dialysis", "transplant"):
            slide = presentation.slides.add_slide(presentation.slide_layouts[5])
            slide.shapes.title.text = "Synthetic " + topic + " study"
            if not title_only:
                slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = "Synthetic " + topic + " access learning observation."
                table = slide.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(6), Inches(2)).table
                table.cell(0, 0).text, table.cell(0, 1).text = "Topic", "Observation"
                for n in range(1, 3):
                    table.cell(n, 0).text, table.cell(n, 1).text = topic, "Synthetic table observation " + str(n)
        presentation.save(path)
    else:
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.remove(workbook.active)
        for topic in ("CKD", "Electrolytes"):
            sheet = workbook.create_sheet(topic)
            sheet.append(["Topic", "Observation"])
            for n in range(70 if long_table else 2):
                sheet.append(["ROW" + str(n), "Synthetic " + topic + " study observation " * (12 if long_table else 1)])
        workbook.save(path)
    return path


@pytest.fixture
def offline_guard(monkeypatch):
    connect, connect_ex = socket.socket.connect, socket.socket.connect_ex
    resolve = socket.getaddrinfo
    def guard(original):
        def wrapped(sock, address):
            if isinstance(address, tuple) and ipaddress.ip_address(address[0]).is_loopback:
                return original(sock, address)
            pytest.fail("External network requested by Office extraction")
        return wrapped
    monkeypatch.setattr(socket.socket, "connect", guard(connect))
    monkeypatch.setattr(socket.socket, "connect_ex", guard(connect_ex))
    def local_resolve(host, *args, **kwargs):
        if host in (None, "localhost"):
            return resolve(host, *args, **kwargs)
        try:
            local = ipaddress.ip_address(host).is_loopback
        except ValueError:
            local = False
        if not local:
            pytest.fail("External name resolution requested by Office extraction")
        return resolve(host, *args, **kwargs)
    monkeypatch.setattr(socket, "getaddrinfo", local_resolve)
    monkeypatch.setattr(subprocess, "Popen", lambda *a, **k: pytest.fail("Office extraction launched an external process"))


@pytest.fixture
def office_extractor(tmp_path, offline_guard):
    from tokenizers import Tokenizer, models, pre_tokenizers
    selected = os.environ.get("RENULUS_OFFICE_TOKENIZER")
    if selected:
        token_path = Path(selected)
        manifest = json.loads((Path(__file__).parents[2] / "packaging/runtime/helper-assets.json").read_text())
        record = next(item for item in manifest["groups"]["embedding"]["files"] if item["path"].endswith("/tokenizer.json"))
        assert hashlib.sha256(token_path.read_bytes()).hexdigest() == record["sha256"]
        root = token_path.parent
    else:
        # Application tests are portable; the proof command selects the pinned BGE JSON.
        root = tmp_path / "synthetic-tokenizer"
        root.mkdir()
        tokenizer = Tokenizer(models.WordLevel({"[UNK]": 0}, unk_token="[UNK]"))
        tokenizer.pre_tokenizer = pre_tokenizers.Whitespace()
        tokenizer.save(str(root / "tokenizer.json"))
    class TokenizerOnlyAssets:
        def validate(self, kind):
            assert kind == "fastembed", "Office requested PDF/OCR model assets"
            return root
    return DoclingExtractor(TokenizerOnlyAssets())


@pytest.mark.parametrize("suffix", OFFICE_MEDIA)
def test_real_office_text_table_and_typed_provenance(office_extractor, tmp_path, suffix):
    path = make_office(tmp_path / ("synthetic" + suffix))
    original = path.read_bytes()
    result = office_extractor.extract_file(path, "Synthetic Office study")
    assert path.read_bytes() == original
    assert result.ocr == {"used": False, "confidence": None}
    assert result.document["tables"]
    assert "Synthetic" in " \n".join(p["context_text"] for p in result.passages)
    locators = [locator for p in result.passages for locator in p["locators"]]
    assert locators and all(p["page"] is None and p["format"] == suffix[1:] for p in locators)
    assert any(p.get("table_ref", "").startswith("#/tables/") for p in locators)
    for locator in locators:
        kind, ordinal = locator["item_ref"].removeprefix("#/").split("/")
        assert result.document[kind][int(ordinal)]["self_ref"] == locator["item_ref"]
    if suffix == ".pptx":
        assert {p["slide"] for p in locators} == {1, 2}
        assert all(p["bbox"] and p["slide_size"] for p in locators)
    elif suffix == ".xlsx":
        assert {(p["sheet"], p["sheet_name"]) for p in locators} == {(1, "CKD"), (2, "Electrolytes")}
        assert all(p["cell_bbox"] for p in locators)
    else:
        assert all("slide" not in p and "sheet" not in p and "bbox" not in p for p in locators)
    for pipeline in office_extractor._office_converter.initialized_pipelines.values():
        assert type(pipeline).__name__ == "SimplePipeline"
        assert all(model.enabled is False for model in pipeline.enrichment_pipe)
        assert pipeline.pipeline_options.enable_remote_services is False
        assert pipeline.pipeline_options.allow_external_plugins is False


@pytest.mark.parametrize("suffix", [".docx", ".xlsx"])
def test_real_office_hybrid_table_split_keeps_all_rows_and_budget(office_extractor, tmp_path, suffix):
    from tokenizers import Tokenizer
    path = make_office(tmp_path / ("long" + suffix), long_table=True)
    extracted = office_extractor.extract_file(path, "Synthetic long table")
    tokenizer = Tokenizer.from_file(str(office_extractor.assets.validate("fastembed") / "tokenizer.json"))
    tokenizer.no_truncation()
    tokenizer.no_padding()
    assert len(extracted.passages) > 2
    assert all(len(tokenizer.encode(p["context_text"], add_special_tokens=False).ids) <= CHUNK_TOKENS for p in extracted.passages)
    text = "\n".join(p["text"] for p in extracted.passages)
    assert set(re.findall(r"ROW[0-9]+\b", text)) == {"ROW" + str(n) for n in range(70)}
    assert all(p["locators"] and all(l["page"] is None for l in p["locators"]) for p in extracted.passages)


@pytest.mark.parametrize("suffix", [".docx", ".pptx"])
def test_title_only_office_keeps_original_docling_structure(office_extractor, tmp_path, suffix):
    path = make_office(tmp_path / ("title-only" + suffix), title_only=True)
    result = office_extractor.extract_file(path, "Synthetic title only")
    assert result.passages and all(p["locators"] for p in result.passages)
    assert any(item["label"] in ("title", "section_header") for item in result.document["texts"])


@pytest.mark.parametrize("suffix", OFFICE_MEDIA)
def test_empty_office_is_still_a_failed_extraction(office_extractor, tmp_path, suffix):
    path = tmp_path / ("empty" + suffix)
    if suffix == ".docx":
        from docx import Document
        Document().save(path)
    elif suffix == ".pptx":
        from pptx import Presentation
        Presentation().save(path)
    else:
        from openpyxl import Workbook
        Workbook().save(path)
    with pytest.raises(ApiError) as error:
        office_extractor.extract_file(path, "Empty synthetic Office input")
    assert error.value.code in ("empty_extraction", "extraction_failed")


def test_real_docx_hyperlink_text_and_xlsx_formula_never_follow_links(office_extractor, tmp_path):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE
    from openpyxl import Workbook
    word = Document()
    paragraph = word.add_paragraph("Synthetic text with an inert link: ")
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), paragraph.part.relate_to("https://example.invalid/never-fetch", RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
    run, text = OxmlElement("w:r"), OxmlElement("w:t")
    text.text = "Synthetic link label"
    run.append(text)
    link.append(run)
    paragraph._p.append(link)
    path = tmp_path / "hyperlink.docx"
    word.save(path)
    extracted = office_extractor.extract_file(path, "Synthetic inert hyperlink")
    assert "Synthetic link label" in "\n".join(p["text"] for p in extracted.passages)
    workbook = Workbook()
    workbook.active.append(["Synthetic", "Safe observation"])
    workbook.active.append(["Formula", '=HYPERLINK("https://example.invalid/never-fetch","Never execute")'])
    path = tmp_path / "formula.xlsx"
    workbook.save(path)
    extracted = office_extractor.extract_file(path, "Synthetic inert formula")
    assert "never-fetch" not in "\n".join(p["text"] for p in extracted.passages)


def test_bad_embedded_word_picture_cannot_trigger_libreoffice_fallback(office_extractor, tmp_path, monkeypatch):
    from docx import Document
    from PIL import Image
    from docling.backend import msword_backend
    monkeypatch.setattr(msword_backend, "get_docx_to_pdf_converter", lambda: pytest.fail("Discovered an external Word renderer"))
    picture = BytesIO()
    Image.new("RGB", (2, 2), "white").save(picture, format="PNG")
    picture.seek(0)
    word = Document()
    word.add_paragraph("Synthetic kidney teaching with a broken image.")
    word.add_picture(picture)
    original = BytesIO()
    word.save(original)
    replacement = BytesIO()
    with ZipFile(original) as source, ZipFile(replacement, "w") as target:
        for item in source.infolist():
            target.writestr(item, b"broken synthetic image" if item.filename.startswith("word/media/") else source.read(item))
    path = tmp_path / "bad-picture.docx"
    path.write_bytes(replacement.getvalue())
    result = office_extractor.extract_file(path, "Synthetic broken picture")
    assert "Synthetic kidney teaching" in "\n".join(p["text"] for p in result.passages)
    assert result.ocr["used"] is False


@pytest.mark.parametrize("suffix", OFFICE_MEDIA)
def test_temporary_office_bytes_still_refused_before_asset_or_file_use(office_extractor, tmp_path, suffix):
    before = list(tmp_path.rglob("*"))
    with pytest.raises(ApiError) as error:
        office_extractor.extract_bytes(b"synthetic", "input" + suffix, "Temporary")
    assert error.value.code == "unsupported_file"
    assert list(tmp_path.rglob("*")) == before
