# SPDX-License-Identifier: MIT
"""Selected pinned parsers to inert Office, PDF and image reader payloads."""
from datetime import date, datetime, time
from contextlib import closing
from io import BytesIO
import math
import re
from threading import RLock
import warnings

from ..contracts import ApiError
from .office import validate_office
from .reader_models import ReaderBlock, ReaderCell, ReaderDiagnostic, ReaderPayload, ReaderTable
from .reader_text import block, table_block

PDF_LOCK = RLock()  # PDFium calls are process-serial, including rasterisation.
MAX_PIXELS = 12_000_000


def _image(data, frame=1):
    from PIL import Image
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            image = Image.open(BytesIO(data))
            if image.format not in {"PNG", "JPEG", "TIFF"}:
                raise ApiError("reader_image_invalid", "Choose PNG, JPEG or TIFF", 422)
            frames = getattr(image, "n_frames", 1)
            if frames > 300 or image.width * image.height > MAX_PIXELS:
                raise ApiError("reader_image_limit", "The image exceeds 300 frames or 12 million pixels", 413)
            if not 1 <= frame <= frames:
                raise ApiError("reader_page_missing", "The requested image frame is unavailable", 404)
            image.seek(frame - 1)
            if image.width * image.height > MAX_PIXELS:
                raise ApiError("reader_image_limit", "The image frame exceeds 12 million pixels", 413)
            image.load()
            output = BytesIO()
            image.convert("RGB").save(output, format="PNG")
            image.close()
            return output.getvalue(), frames
    except ApiError:
        raise
    except Exception:
        raise ApiError("reader_image_invalid", "The image could not be decoded safely", 422) from None


def image_payload(data, request):
    _, frames = _image(data, request.page)
    return ReaderPayload(kind="image", page=request.page, pages=frames,
        asset_id=f"image:{request.page}", diagnostics=(ReaderDiagnostic(code="ocr_separate",
        message="This preview does not establish OCR or model image interpretation."),))


def pdf_payload(data, request):
    import pypdfium2 as pdfium
    try:
        with PDF_LOCK, pdfium.PdfDocument(data) as document:
            pages = len(document)
            if pages > 300:
                raise ApiError("reader_page_limit", "The PDF exceeds 300 pages", 413)
            if not 1 <= request.page <= pages:
                raise ApiError("reader_page_missing", "The requested physical page is unavailable", 404)
            with closing(document[request.page - 1]) as page, closing(page.get_textpage()) as textpage:
                text = textpage.get_text_range().replace("\r\n", "\n")
            if len(text) > 200000:
                raise ApiError("reader_text_limit", "The PDF page contains too much text", 413)
            ref = f"reader-v1:pdf:{request.page}"
            blocks = (ReaderBlock(kind="paragraph", item_ref=ref, text=text,
                locator={"kind": "pdf", "physical_page": request.page}),) if text.strip() else ()
            diagnostics = () if blocks else (ReaderDiagnostic(code="pdf_no_text",
                message="This page has no embedded text. OCR has not been performed by this reader."),)
            return ReaderPayload(kind="pdf", blocks=blocks, page=request.page, pages=pages,
                asset_id=f"pdf:{request.page}", diagnostics=diagnostics)
    except ApiError:
        raise
    except Exception:
        raise ApiError("reader_pdf_invalid", "The PDF is incomplete, encrypted or could not be parsed", 422) from None


def _display_cell(value, number_format):
    if value is None:
        return "", None, False
    if isinstance(value, (datetime, date, time)):
        return value.isoformat(), value.isoformat(), True
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE", value, False
    if isinstance(value, (int, float)):
        if not math.isfinite(value):
            raise ApiError("reader_cell_invalid", "A spreadsheet cell contains a non-finite value", 422)
        if number_format in {"General", "@"}:
            return str(value), value, False
        percentage = re.fullmatch(r"0(?:\.(0+))?%", number_format)
        if percentage:
            digits = len(percentage[1] or "")
            return f"{value * 100:.{digits}f}%", value, False
        if re.fullmatch(r"0{2,12}", number_format) and float(value).is_integer():
            return f"{int(value):0{len(number_format)}d}", value, False
        decimal = re.fullmatch(r"(#,##)?0(?:\.(0+))?", number_format)
        if decimal:
            digits = len(decimal[2] or "")
            return format(value, f",.{digits}f" if decimal[1] else f".{digits}f"), value, False
        return str(value), value, True
    return str(value), str(value), False


def xlsx_payload(data, request):
    from openpyxl import load_workbook
    validate_office(data, ".xlsx", max_bytes=64 * 1024 * 1024, max_pages=300)
    formulas = values = None
    try:
        formulas = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        values = load_workbook(BytesIO(data), read_only=True, data_only=True, keep_links=False)
        sheets = formulas.sheetnames
        sheet = request.sheet or sheets[0]
        if sheet not in sheets:
            raise ApiError("reader_sheet_missing", "The requested sheet is unavailable", 404)
        source, cached = formulas[sheet], values[sheet]
        if source.max_row * source.max_column > 100000 or source.max_column > 200:
            raise ApiError("reader_table_limit", "The sheet exceeds 100,000 cells or 200 columns", 413)
        rows, diagnostics = [], {}
        end = min(source.max_row, request.first_row + request.row_limit - 1)
        if request.first_row <= end:
            source_rows = source.iter_rows(min_row=request.first_row, max_row=end)
            cached_rows = cached.iter_rows(min_row=request.first_row, max_row=end)
            for source_row, cached_row in zip(source_rows, cached_rows, strict=True):
                cells = []
                for cell, cache in zip(source_row, cached_row, strict=True):
                    formula = str(cell.value) if cell.data_type == "f" else None
                    value = cache.value if formula else cell.value
                    display, literal, uncertain = _display_cell(value, cell.number_format)
                    if formula and value is None:
                        display = "Cached value unavailable"
                        diagnostics["formula_cache_missing"] = "Formula results without cached values are unavailable; formulas are never executed."
                    if uncertain:
                        diagnostics["number_format_partial"] = "Some Excel formats are unsupported; their literal values and format strings are retained."
                    if len(display) > 20000 or (formula and len(formula) > 20000):
                        raise ApiError("reader_table_limit", "A sheet cell exceeds 20,000 characters", 413)
                    cells.append(ReaderCell(display=display, value=literal, formula=formula,
                        cached=not formula or value is not None, number_format=cell.number_format))
                rows.append(tuple(cells))
        diagnostics["xlsx_layout"] = "Sheets use a row/column grid; charts, drawings and merged-cell layout are not reproduced."
        return ReaderPayload(kind="table", table=ReaderTable(format="xlsx", sheet_name=sheet,
            sheets=tuple(sheets), first_row=request.first_row, total_rows=source.max_row,
            total_columns=source.max_column, rows=tuple(rows)), diagnostics=tuple(
                ReaderDiagnostic(code=code, message=message) for code, message in diagnostics.items()))
    finally:
        if formulas:
            formulas.close()
        if values:
            values.close()


def docx_payload(data, request):
    from docx import Document
    from docx.table import Table
    validate_office(data, ".docx", max_bytes=64 * 1024 * 1024, max_pages=300)
    document = Document(BytesIO(data))
    blocks, characters = [], 0
    for item in document.iter_inner_content():
        if isinstance(item, Table):
            if item.rows:
                if any(cell._tc.xpath("./w:tcPr/w:gridSpan | ./w:tcPr/w:vMerge")
                       for row in item.rows for cell in row.cells):
                    raise ApiError("reader_table_unsupported", "Merged DOCX table cells are not supported by this reader", 422)
                header = bool(item.rows[0]._tr.xpath("./w:trPr/w:tblHeader"))
                blocks.append(table_block([[cell.text for cell in row.cells] for row in item.rows],
                    len(blocks), header_rows=int(header)))
        else:
            text = item.text
            style = item.style.name if item.style else ""
            heading = re.fullmatch(r"Heading ([1-6])", style)
            if text:
                kind = "heading" if heading else "list_item" if item._p.xpath("./w:pPr/w:numPr") else "paragraph"
                blocks.append(block(kind, text, len(blocks), level=int(heading[1]) if heading else None,
                                    section=style[:240] or None))
            for image in item._p.xpath(".//a:blip"):
                rel = image.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
                if rel and rel in document.part.rels and not document.part.rels[rel].is_external:
                    blocks.append(ReaderBlock(kind="image", item_ref=f"reader-v1:image:{len(blocks)}",
                                              asset_id=f"docx:{rel}"))
        characters = sum(len(item.text) for item in blocks)
        if len(blocks) > 4000 or characters > 2_000_000:
            raise ApiError("reader_text_limit", "The document exceeds its text or block limit", 413)
    return ReaderPayload(kind="text", blocks=tuple(blocks), diagnostics=(ReaderDiagnostic(
        code="docx_reflow", message="Document content is reflowed; sections and text locations do not imply original page numbers."),))


def binary_asset(data, suffix, asset_id):
    if suffix == ".pdf" and re.fullmatch(r"pdf:[1-9][0-9]{0,2}", asset_id):
        import pypdfium2 as pdfium
        try:
            with PDF_LOCK, pdfium.PdfDocument(data) as document:
                number = int(asset_id.split(":")[1])
                if len(document) > 300 or number > len(document):
                    raise ApiError("reader_page_missing", "The physical page is unavailable", 404)
                with closing(document[number - 1]) as page:
                    width, height = page.get_size()
                    scale = min(1.5, math.sqrt(MAX_PIXELS / max(1, width * height)))
                    bitmap = page.render(scale=scale)
                    try:
                        output = BytesIO()
                        with bitmap.to_pil() as image:
                            image.save(output, format="PNG")
                        return output.getvalue()
                    finally:
                        bitmap.close()
        except ApiError:
            raise
        except Exception:
            raise ApiError("reader_pdf_invalid", "The PDF page could not be rendered", 422) from None
    if suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff"} and re.fullmatch(r"image:[1-9][0-9]{0,2}", asset_id):
        return _image(data, int(asset_id.split(":")[1]))[0]
    if suffix == ".docx" and re.fullmatch(r"docx:rId[0-9]{1,8}", asset_id):
        from docx import Document
        validate_office(data, suffix, max_bytes=64 * 1024 * 1024, max_pages=300)
        document = Document(BytesIO(data))
        relation = document.part.rels.get(asset_id.split(":")[1])
        if relation and not relation.is_external and relation.reltype.endswith("/image"):
            return _image(relation.target_part.blob)[0]
    raise ApiError("reader_asset_missing", "This asset is not part of the selected original", 404)
